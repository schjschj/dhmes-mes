# -*- coding: utf-8 -*-
"""작업지시서 (주/야간 생산계획) — 하루 한 장, 양식 내용은 JSON.

data = {left:[line], right:[line], t60:[[{die,code}×3]], elec:[{die,code}], etc:cell, prep:cell}
line = {equip, label, star, cells:{cur, die, code, time, next}}
cell = {t, fill, font, r?}   (r = 글자색 여러 가지 [[글자, 'RRGGBB'|''], ...])
"""
import datetime
import io
import json
import re

import db

WD = "월화수목금토일"
CUR_FILL = "B7DEE8"
STOP_FILLS = ("FF0000", "0000FF")
CODE_RE = re.compile(r"[0-9A-Z]{3,}-[0-9A-Z]+", re.I)


def day_label(date):
    d = datetime.date.fromisoformat(date)
    return "%s (%s)" % (date, WD[d.weekday()])


def cell(t="", fill="", font=""):
    return {"t": t, "fill": fill, "font": font}


def _doc(row):
    if not row:
        return None
    d = dict(row)
    d["data"] = json.loads(d["data"] or "{}")
    d.pop("draft", None)
    return d


def get(date):
    row = db.one("SELECT * FROM work_orders WHERE date=?", (date,))
    prev = db.one("SELECT date FROM work_orders WHERE date<? ORDER BY date DESC LIMIT 1", (date,))
    return {"doc": _doc(row), "day": day_label(date), "prev": prev["date"] if prev else None}


def list_all():
    return db.rows("SELECT id, date, status, written_by FROM work_orders ORDER BY date DESC")


def _last_codes(before=None):
    """설비별 최근 생산 제번."""
    sql = "SELECT equip_no, code, date FROM production"
    args = []
    if before:
        sql += " WHERE date <= ?"
        args.append(before)
    sql += " ORDER BY date, id"
    out = {}
    for r in db.rows(sql, args):
        out[str(r["equip_no"])] = r["code"]
    return out


def blank(date):
    eqs = db.rows("SELECT no, ton FROM equipment WHERE COALESCE(status,'') <> '폐기'")
    eqs.sort(key=lambda e: (len(str(e["no"])), str(e["no"])))
    last = _last_codes(date)
    lines = []
    for e in eqs:
        no = str(e["no"])
        label = no + "호기" + ("\n(%sT)" % int(e["ton"]) if e.get("ton") else "")
        lines.append({"equip": no, "label": label, "star": False,
                      "cells": {"cur": cell(last.get(no, ""), CUR_FILL), "die": cell("☞☞"), "code": cell("☞☞"),
                                "time": cell(), "next": cell()}})
    half = (len(lines) + 1) // 2
    return {"left": lines[:half], "right": lines[half:],
            "t60": [[{"die": cell(), "code": cell()} for _ in range(3)] for _ in range(6)],
            "elec": [{"die": cell(), "code": cell()} for _ in range(6)],
            "etc": cell(), "prep": cell()}


def next_day(data):
    """직전 작업지시서 → 다음날. 교체 제번은 현 생산품으로, 노랑(당일)·완료 표시는 지움."""
    data = json.loads(json.dumps(data))
    moved = []
    for side in ("left", "right"):
        for ln in data.get(side, []):
            c = ln["cells"]
            codes = CODE_RE.findall(c.get("code", {}).get("t") or "")
            if codes:
                c["cur"] = cell("\n".join(codes), CUR_FILL)
                moved.append("%s호기 %s" % (ln.get("equip"), ",".join(codes)))
                c["code"] = cell("☞☞")
            die = c.get("die") or cell()
            stop = die.get("fill") in STOP_FILLS or die.get("font") == "0000FF"
            if not stop and (die.get("fill") == "FFFF00" or (die.get("t") or "").split("\n")[0] in ("생산완료", "출근즉시", "재료소진")):
                c["die"] = cell("☞☞")
            if (c.get("time") or {}).get("fill") == "FFFF00":
                c["time"] = cell()
            nx = c.get("next") or cell()
            if nx.get("fill") == "92D050":
                nx["fill"] = "FFFF00"
            elif nx.get("fill") == "B3A2C7":
                nx["fill"] = "92D050"
            c["next"] = nx
            if c["cur"].get("fill") in STOP_FILLS and not stop:
                c["cur"]["fill"], c["cur"]["font"] = CUR_FILL, ""
    return data, moved


def create(date, from_date, mode, user):
    if db.one("SELECT id FROM work_orders WHERE date=?", (date,)):
        raise ValueError("이미 작업지시서가 있습니다.")
    moved = []
    if from_date:
        src = db.one("SELECT data FROM work_orders WHERE date=?", (from_date,))
        if not src:
            raise ValueError("%s 작업지시서가 없습니다." % from_date)
        data = json.loads(src["data"])
        if mode == "next":
            data, moved = next_day(data)
    else:
        data = blank(date)
    now = db.now()
    cur = db.run("""INSERT INTO work_orders(date,data,version,status,written_by,written_at,created,created_by,updated,updated_by)
                    VALUES(?,?,1,'작성',?,?,?,?,?,?)""",
                 (date, json.dumps(data, ensure_ascii=False), user, now, now, user, now, user))
    return {"id": cur.lastrowid, "moved": moved}


def save(wid, data, version, user):
    row = db.one("SELECT * FROM work_orders WHERE id=?", (wid,))
    if not row:
        raise ValueError("작업지시서를 찾을 수 없습니다.")
    if version is not None and int(version) != int(row["version"]):
        raise ValueError("다른 PC 에서 먼저 고쳤습니다. 새로고침(F5) 한 뒤 다시 고쳐 주세요.")
    reset = row["status"] not in (None, "", "작성")
    now = db.now()
    db.run("""UPDATE work_orders SET data=?, version=version+1, written_by=?, written_at=?, updated=?, updated_by=?
              %s WHERE id=?""" % (", status='작성', reviewed_by=NULL, reviewed_at=NULL, approved_by=NULL, approved_at=NULL" if reset else ""),
           (json.dumps(data, ensure_ascii=False), user, now, now, user, wid))
    return {"version": row["version"] + 1, "reset": reset}


def sign(wid, kind, user):
    row = db.one("SELECT * FROM work_orders WHERE id=?", (wid,))
    if not row:
        raise ValueError("작업지시서를 찾을 수 없습니다.")
    now = db.now()
    if kind == "review":
        db.run("UPDATE work_orders SET status='검토', reviewed_by=?, reviewed_at=? WHERE id=?", (user, now, wid))
    elif kind == "approve":
        db.run("UPDATE work_orders SET status='승인', approved_by=?, approved_at=?, reviewed_by=COALESCE(reviewed_by, ?), reviewed_at=COALESCE(reviewed_at, ?) WHERE id=?",
               (user, now, user, now, wid))
        try:
            import plan
            plan.learn_from_wo(row["date"])
        except Exception:
            pass
    elif kind == "cancel":
        db.run("UPDATE work_orders SET status='작성', reviewed_by=NULL, reviewed_at=NULL, approved_by=NULL, approved_at=NULL WHERE id=?", (wid,))
    else:
        raise ValueError("알 수 없는 결재입니다.")
    return row["date"]


def delete(wid):
    row = db.one("SELECT date FROM work_orders WHERE id=?", (wid,))
    db.run("DELETE FROM work_orders WHERE id=?", (wid,))
    return row["date"] if row else ""


def hint(date):
    """설비별 직전 생산일의 실적 (작업지시서와 비교용)."""
    out = {}
    for r in db.rows("""SELECT equip_no, code, date, COALESCE(day_reels,0)+COALESCE(night_reels,0) reels FROM production
                        WHERE date < ? ORDER BY date DESC""", (date,)):
        eq = str(r["equip_no"])
        lst = out.setdefault(eq, [])
        if lst and lst[0]["date"] != r["date"]:
            continue
        lst.append({"date": r["date"], "code": r["code"], "reels": r["reels"]})
    return out


def summary(date):
    row = db.one("SELECT * FROM work_orders WHERE date=?", (date,))
    if not row:
        return {"exists": False}
    d = json.loads(row["data"])
    change, stop, night = [], [], 0
    for side in ("left", "right"):
        for ln in d.get(side, []):
            c = ln["cells"]
            if ln.get("star"):
                night += 1
            code_t = (c.get("code") or {}).get("t") or ""
            if code_t.strip() and code_t.strip() != "☞☞":
                change.append({"equip": ln["label"].split("\n")[0], "code": code_t.split("\n")[0],
                               "die": (c.get("die") or {}).get("t", "").split("\n")[0] if (c.get("die") or {}).get("t") != "☞☞" else "",
                               "time": (c.get("time") or {}).get("t", "")})
            die = c.get("die") or {}
            if die.get("fill") in STOP_FILLS:
                stop.append({"equip": ln["label"].split("\n")[0], "why": (die.get("t") or "").split("\n")[0]})
    return {"exists": True, "status": row["status"], "written_by": row["written_by"], "reviewed_by": row["reviewed_by"],
            "approved_by": row["approved_by"], "change": change, "stop": stop, "night": night}


# ──────────────────────────────────────────────── 엑셀
def to_xlsx(wid):
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    row = db.one("SELECT * FROM work_orders WHERE id=?", (wid,))
    d = json.loads(row["data"])
    wb = Workbook()
    ws = wb.active
    ws.title = row["date"]
    thin = Side(style="thin", color="000000")
    bd = Border(left=thin, right=thin, top=thin, bottom=thin)

    def put(r, c, cl, bold=False):
        cl = cl or {}
        x = ws.cell(row=r, column=c, value=cl.get("t") if isinstance(cl, dict) else cl)
        if isinstance(cl, dict):
            if cl.get("fill"):
                x.fill = PatternFill("solid", fgColor=cl["fill"])
            x.font = Font(color=cl.get("font") or "000000", bold=bold)
        x.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        x.border = bd

    ws.cell(row=1, column=1, value=day_label(row["date"]))
    ws.cell(row=1, column=4, value="주/야간 생산계획")
    ws.cell(row=1, column=11, value=row["status"])
    heads = ["설비", "제번", "Die No", "제번", "교체시간", "차기생산품"] * 2
    for i, h in enumerate(heads, 1):
        put(3, i, h, True)
    n = max(len(d.get("left", [])), len(d.get("right", [])))
    for i in range(n):
        for off, side in ((0, "left"), (6, "right")):
            ln = (d.get(side) or [None] * n)[i] if i < len(d.get(side, [])) else None
            if not ln:
                continue
            put(4 + i, off + 1, ("★" if ln.get("star") else "") + ln.get("label", ""), True)
            for j, k in enumerate(("cur", "die", "code", "time", "next")):
                put(4 + i, off + 2 + j, ln["cells"].get(k))
    r0 = 5 + n
    put(r0, 1, "60톤", True)
    put(r0, 7, "전자용", True)
    for j, rowp in enumerate(d.get("t60", [])):
        for k, pr in enumerate(rowp):
            put(r0 + 1 + j, 1 + k * 2, pr.get("die"))
            put(r0 + 1 + j, 2 + k * 2, pr.get("code"))
    for j, e in enumerate(d.get("elec", [])):
        put(r0 + 1 + j, 7, e.get("die"))
        put(r0 + 1 + j, 8, e.get("code"))
    put(r0, 9, "기타 전달사항", True)
    put(r0 + 1, 9, d.get("etc"))
    put(r0 + 4, 9, "금형 준비", True)
    put(r0 + 5, 9, d.get("prep"))
    for col in "ABCDEFGHIJKL":
        ws.column_dimensions[col].width = 14
    bio = io.BytesIO()
    wb.save(bio)
    return row["date"], bio.getvalue()


def import_xlsx(buf, user):
    """예전 엑셀 작업지시서 (시트마다 하루). 시트 이름이나 첫 칸에서 날짜를 찾고 위 양식 순서로 읽음."""
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(buf), data_only=True)
    report = []
    for ws in wb.worksheets:
        date = None
        for txt in [ws.title] + [str(ws.cell(row=r, column=c).value or "") for r in range(1, 4) for c in range(1, 4)]:
            m = re.search(r"(20\d\d)[.\-/년 ]+(\d{1,2})[.\-/월 ]+(\d{1,2})", txt)
            if m:
                date = "%s-%02d-%02d" % (m.group(1), int(m.group(2)), int(m.group(3)))
                break
            m = re.search(r"^(\d{1,2})[.\-/월 ]+(\d{1,2})", txt.strip())
            if m:
                date = "%d-%02d-%02d" % (datetime.date.today().year, int(m.group(1)), int(m.group(2)))
                break
        if not date:
            report.append("%s: 날짜를 찾지 못해 건너뜀" % ws.title)
            continue
        if db.one("SELECT id FROM work_orders WHERE date=?", (date,)):
            report.append("%s: 이미 있어서 건너뜀" % date)
            continue
        data = blank(date)
        lines = {ln["equip"]: ln for ln in data["left"] + data["right"]}
        hit = 0
        for r in range(1, ws.max_row + 1):
            for off in (0, 6):
                lab = str(ws.cell(row=r, column=off + 1).value or "")
                m = re.match(r"\s*★?\s*(?:DH-)?(\d+)\s*호?", lab)
                if not m or m.group(1) not in lines:
                    continue
                ln = lines[m.group(1)]
                ln["star"] = "★" in lab
                for j, k in enumerate(("cur", "die", "code", "time", "next")):
                    x = ws.cell(row=r, column=off + 2 + j)
                    fill = ""
                    try:
                        if x.fill and x.fill.fgColor and x.fill.fgColor.type == "rgb" and x.fill.fill_type:
                            fill = str(x.fill.fgColor.rgb)[-6:]
                    except Exception:
                        fill = ""
                    ln["cells"][k] = cell(str(x.value) if x.value is not None else "", "" if fill in ("000000", "FFFFFF") else fill)
                hit += 1
        now = db.now()
        db.run("""INSERT INTO work_orders(date,data,version,status,written_by,written_at,created,created_by,updated,updated_by)
                  VALUES(?,?,1,'작성',?,?,?,?,?,?)""", (date, json.dumps(data, ensure_ascii=False), user, now, now, user, now, user))
        report.append("%s: 가져옴 (설비 %d줄)" % (date, hit))
    return report
