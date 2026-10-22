# -*- coding: utf-8 -*-
"""시스템: 사용자·권한 · 설정 · 백업 · 프로그램 업데이트 · 메뉴 배치 · 데이터 가져오기 · 한줄 메모 · 접속 프로그램."""
import datetime
import glob
import io
import json
import os
import re
import shutil
import sqlite3
import sys
import threading
import time
import zipfile

import db
import meta
from common import get_json, set_json

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BACKUP_DIR = os.path.join(db.DATA_DIR, "backup")
UPD_DIR = os.path.join(db.DATA_DIR, "update")
KEEP_DAYS = 60
LEVELS = {0: "숨김", 1: "보기", 2: "입력·수정", 3: "삭제까지"}
SETTING_KEYS = ["company", "mold_alarm1", "mold_alarm2", "target_month_qty", "target_defect_rate", "target_oee", "interlock_first"]


# ──────────────────────────────────────────── 사용자 · 권한
def _groups():
    gs = meta.effective(get_json("menu_layout", None))
    return gs


def users_view():
    gs = _groups()
    out = []
    for u in db.rows("SELECT * FROM users ORDER BY admin DESC, login"):
        out.append({k: u[k] for k in ("id", "login", "name", "admin", "active", "must_change", "note", "last_login",
                                     "created", "created_by", "updated", "updated_by")} | {"perms": json.loads(u["perms"] or "{}")})
    return {"users": out, "groups": [g["key"] for g in gs], "levels": LEVELS,
            "group_names": {g["key"]: g["name"] for g in gs},
            "group_menus": {g["key"]: [it["label"] for it in g["items"]] for g in gs}}


def user_save(uid, b, by):
    login, name = (b.get("login") or "").strip(), (b.get("name") or "").strip()
    if not login or not name:
        raise ValueError("아이디와 이름을 넣으세요.")
    other = db.one("SELECT id FROM users WHERE login=? AND id<>?", (login, uid or 0))
    if other:
        raise ValueError("같은 아이디가 이미 있습니다.")
    perms = {k: int(v) for k, v in (b.get("perms") or {}).items()}
    admin = 1 if b.get("admin") else 0
    active = 1 if b.get("active") else 0
    if uid:
        old = db.one("SELECT * FROM users WHERE id=?", (uid,))
        if old["admin"] and (not admin or not active) and db.one("SELECT COUNT(*) n FROM users WHERE admin=1 AND active=1 AND id<>?", (uid,))["n"] == 0:
            raise ValueError("관리자가 한 명은 있어야 합니다.")
        db.run("UPDATE users SET login=?, name=?, admin=?, active=?, perms=?, note=?, updated=?, updated_by=? WHERE id=?",
               (login, name, admin, active, json.dumps(perms), b.get("note") or "", db.now(), by, uid))
        if b.get("password"):
            if len(b["password"]) < 4:
                raise ValueError("비밀번호는 4자 이상이어야 합니다.")
            db.set_password(uid, b["password"], by)
            db.run("UPDATE users SET must_change=1 WHERE id=?", (uid,))
        return {"id": uid}
    if len(b.get("password") or "") < 4:
        raise ValueError("처음 비밀번호를 4자 이상 넣으세요.")
    u = db.create_user(login, name, b["password"], admin=admin, perms=perms, by=by, must_change=1)
    db.run("UPDATE users SET active=?, note=? WHERE id=?", (active, b.get("note") or "", u["id"]))
    return {"id": u["id"]}


def user_delete(uid, me_id):
    if uid == me_id:
        raise ValueError("자기 계정은 지울 수 없습니다.")
    u = db.one("SELECT * FROM users WHERE id=?", (uid,))
    if u and u["admin"] and db.one("SELECT COUNT(*) n FROM users WHERE admin=1 AND id<>?", (uid,))["n"] == 0:
        raise ValueError("마지막 관리자는 지울 수 없습니다.")
    db.run("DELETE FROM sessions WHERE user_id=?", (uid,))
    db.run("DELETE FROM users WHERE id=?", (uid,))


# ──────────────────────────────────────────── 설정
def settings_get():
    return {k: db.get_setting(k) for k in SETTING_KEYS}


def settings_put(b):
    for k in SETTING_KEYS:
        if k in b:
            db.set_setting(k, "" if b[k] is None else str(b[k]).strip())


# ──────────────────────────────────────────── 백업
def backup_now(tag=""):
    os.makedirs(BACKUP_DIR, exist_ok=True)
    name = "mes_%s%s.db" % (datetime.datetime.now().strftime("%Y%m%d_%H%M%S"), ("_" + tag) if tag else "")
    dst = sqlite3.connect(os.path.join(BACKUP_DIR, name))
    with dst:
        db.conn().backup(dst)
    dst.close()
    cut = time.time() - KEEP_DAYS * 86400
    for f in glob.glob(os.path.join(BACKUP_DIR, "mes_*.db")):
        if os.path.getmtime(f) < cut:
            os.remove(f)
    return name


def backup_list():
    files = []
    for f in sorted(glob.glob(os.path.join(BACKUP_DIR, "mes_*.db")), reverse=True):
        files.append({"name": os.path.basename(f), "size": os.path.getsize(f),
                      "time": datetime.datetime.fromtimestamp(os.path.getmtime(f)).strftime("%Y-%m-%d %H:%M")})
    return {"dir": BACKUP_DIR, "files": files}


def backup_path(name):
    if not name:
        tmp = os.path.join(BACKUP_DIR, "_current.db")
        os.makedirs(BACKUP_DIR, exist_ok=True)
        if os.path.exists(tmp):
            os.remove(tmp)
        dst = sqlite3.connect(tmp)
        with dst:
            db.conn().backup(dst)
        dst.close()
        return tmp, "mes_%s.db" % datetime.datetime.now().strftime("%Y%m%d_%H%M")
    p = os.path.join(BACKUP_DIR, os.path.basename(name))
    return (p, os.path.basename(name)) if os.path.isfile(p) else (None, None)


def daily_backup_loop():
    """하루 한 번 자동 백업 (서버가 켜져 있는 동안)."""
    def run():
        while True:
            try:
                today = datetime.date.today().strftime("%Y%m%d")
                if not glob.glob(os.path.join(BACKUP_DIR, "mes_%s_*.db" % today)):
                    backup_now("auto")
                _check_update_plan()
            except Exception:
                pass
            time.sleep(30)
    threading.Thread(target=run, daemon=True).start()


# ──────────────────────────────────────────── 프로그램 업데이트 (관리자)
def _ver_of(name):
    m = re.search(r"v?(\d+\.\d+(?:\.\d+)?)", name or "")
    return m.group(1) if m else "?"


def update_status():
    staged = get_json("upd_staged", None)
    if staged:
        staged["same"] = staged["version"] == meta.VERSION
        staged["older"] = _vt(staged["version"]) < _vt(meta.VERSION)
    olds = []
    for f in sorted(glob.glob(os.path.join(UPD_DIR, "old_*.zip")), reverse=True)[:3]:
        olds.append({"name": os.path.basename(f), "version": _ver_of(os.path.basename(f)[4:]),
                     "time": datetime.datetime.fromtimestamp(os.path.getmtime(f)).strftime("%Y-%m-%d %H:%M")})
    log = ""
    lp = os.path.join(UPD_DIR, "update.log")
    if os.path.isfile(lp):
        with open(lp, encoding="utf-8", errors="replace") as f:
            log = f.read()[-6000:]
    return {"version": meta.VERSION, "staged": staged, "plan": get_json("upd_plan", None), "result": get_json("upd_result", None),
            "olds": olds, "log": log}


def _vt(v):
    try:
        return tuple(int(x) for x in v.split("."))
    except ValueError:
        return (0,)


def update_upload(name, buf, by):
    try:
        z = zipfile.ZipFile(io.BytesIO(buf))
    except zipfile.BadZipFile:
        raise ValueError("zip 파일이 아닙니다.")
    names = z.namelist()
    if not any(re.search(r"(^|/)server/app\.py$", n) for n in names):
        raise ValueError("이 프로그램의 업데이트 파일이 아닙니다 (server/app.py 가 없음).")
    os.makedirs(UPD_DIR, exist_ok=True)
    with open(os.path.join(UPD_DIR, "staged.zip"), "wb") as f:
        f.write(buf)
    ver = _ver_of(name)
    for n in names:
        if n.endswith("server/meta.py"):
            m = re.search(r'VERSION\s*=\s*"([^"]+)"', z.read(n).decode("utf-8", "replace"))
            if m:
                ver = m.group(1)
    set_json("upd_staged", {"version": ver, "name": name, "size": len(buf), "by": by, "at": db.now()[:16]})
    return {"version": ver}


def update_plan(b, by):
    at = b.get("at") or "now"
    ts = time.time() + 30 if at == "now" else time.mktime(datetime.datetime.strptime(at, "%Y-%m-%d %H:%M").timetuple())
    if b.get("kind") == "rollback":
        name = os.path.basename(b.get("name") or "")
        if not os.path.isfile(os.path.join(UPD_DIR, name)):
            raise ValueError("되돌릴 버전 파일이 없습니다.")
        to_ver = _ver_of(name[4:])
    else:
        st = get_json("upd_staged", None)
        if not st:
            raise ValueError("올린 업데이트 파일이 없습니다.")
        name, to_ver = "staged.zip", st["version"]
    set_json("upd_plan", {"kind": b.get("kind") or "update", "name": name, "from_ver": meta.VERSION, "to_ver": to_ver,
                          "at": datetime.datetime.fromtimestamp(ts).strftime("%Y-%m-%d %H:%M:%S"), "at_ts": ts, "by": by})


def update_cancel():
    db.set_setting("upd_plan", "null")


def update_notice():
    p = get_json("upd_plan", None)
    if not p:
        return None
    secs = int(p["at_ts"] - time.time())
    return {"secs": max(0, secs)} if secs <= 300 else None


def _log_upd(msg):
    os.makedirs(UPD_DIR, exist_ok=True)
    with open(os.path.join(UPD_DIR, "update.log"), "a", encoding="utf-8") as f:
        f.write("[%s] %s\n" % (db.now(), msg))


def _check_update_plan():
    p = get_json("upd_plan", None)
    if not p or time.time() < p["at_ts"]:
        return
    db.set_setting("upd_plan", "null")
    try:
        backup_now("before_update")
        os.makedirs(UPD_DIR, exist_ok=True)
        old = os.path.join(UPD_DIR, "old_v%s_%s.zip" % (meta.VERSION, datetime.datetime.now().strftime("%Y%m%d%H%M")))
        with zipfile.ZipFile(old, "w", zipfile.ZIP_DEFLATED) as z:
            for sub in ("server", "web"):
                for dp, _, fs in os.walk(os.path.join(ROOT, sub)):
                    if "__pycache__" in dp:
                        continue
                    for fn in fs:
                        full = os.path.join(dp, fn)
                        z.write(full, os.path.relpath(full, ROOT))
        with zipfile.ZipFile(os.path.join(UPD_DIR, p["name"])) as z:
            prefix = ""
            for n in z.namelist():
                m = re.match(r"^(.*?)server/app\.py$", n)
                if m:
                    prefix = m.group(1)
                    break
            for n in z.namelist():
                if not n.startswith(prefix) or n.endswith("/"):
                    continue
                rel = n[len(prefix):]
                if not (rel.startswith("server/") or rel.startswith("web/") or rel in ("run.bat", "README.md")):
                    continue
                dst = os.path.join(ROOT, rel)
                os.makedirs(os.path.dirname(dst), exist_ok=True)
                with open(dst, "wb") as f:
                    f.write(z.read(n))
        set_json("upd_result", {"ok": True, "at": db.now()[:16], "msg": "v%s → v%s" % (p["from_ver"], p["to_ver"])})
        _log_upd("업데이트 완료 v%s → v%s, 다시 시작" % (p["from_ver"], p["to_ver"]))
        if p["kind"] != "rollback":
            db.set_setting("upd_staged", "null")
        olds = sorted(glob.glob(os.path.join(UPD_DIR, "old_*.zip")), reverse=True)
        for f in olds[3:]:
            os.remove(f)
        os.execv(sys.executable, [sys.executable] + sys.argv)
    except Exception as e:
        set_json("upd_result", {"ok": False, "at": db.now()[:16], "msg": str(e)})
        _log_upd("업데이트 실패: %s" % e)


# ──────────────────────────────────────────── 메뉴 배치
def menu_layout_view():
    return {"groups": [{"group": g["key"], "name": g["name"], "orig_name": g["orig_name"], "custom": g["custom"],
                        "items": g["items"]} for g in meta.effective(get_json("menu_layout", None))]}


def menu_layout_save(b):
    set_json("menu_layout", {"groups": b.get("groups") or [], "tree": b.get("tree") or [], "labels": b.get("labels") or {}})
    return menu_layout_view()


def menu_layout_reset():
    db.run("DELETE FROM settings WHERE k='menu_layout'")
    return menu_layout_view()


# ──────────────────────────────────────────── 한줄 메모
def memos(me, limit=60, all_=False, q=""):
    total = db.one("SELECT COUNT(*) n FROM memos")["n"]
    sql = "SELECT m.*, (SELECT COUNT(*) FROM memo_comments c WHERE c.memo_id=m.id) n FROM memos m"
    args = []
    if q:
        sql += " WHERE m.text LIKE ? OR m.created_by LIKE ? OR m.id IN (SELECT memo_id FROM memo_comments WHERE text LIKE ?)"
        args = ["%" + q + "%"] * 3
    sql += " ORDER BY m.id DESC"
    if not all_:
        sql += " LIMIT %d" % int(limit)
    rows = db.rows(sql, args)
    for r in rows:
        r["mine"] = r["user_id"] == me["id"] or me.get("admin")
        if all_:
            r["comments"] = _comments(r["id"], me)
    return {"rows": rows, "total": total}


def _comments(mid, me):
    cs = db.rows("SELECT * FROM memo_comments WHERE memo_id=? ORDER BY id", (mid,))
    for c in cs:
        c["mine"] = c["user_id"] == me["id"] or me.get("admin")
    return cs


def memo_one(mid, me):
    m = db.one("SELECT * FROM memos WHERE id=?", (mid,))
    if not m:
        raise ValueError("메모가 없습니다.")
    m["mine"] = m["user_id"] == me["id"] or me.get("admin")
    m["comments"] = _comments(mid, me)
    return m


# ──────────────────────────────────────────── 접속 프로그램 (바탕화면 아이콘)
def client_zip(base_url):
    bat = ("@echo off\r\nset URL=%s/?app=1\r\n"
           "where msedge >nul 2>nul && (start \"\" msedge --app=%%URL%% & exit /b)\r\n"
           "if exist \"%%ProgramFiles(x86)%%\\Microsoft\\Edge\\Application\\msedge.exe\" (start \"\" \"%%ProgramFiles(x86)%%\\Microsoft\\Edge\\Application\\msedge.exe\" --app=%%URL%% & exit /b)\r\n"
           "if exist \"%%ProgramFiles%%\\Google\\Chrome\\Application\\chrome.exe\" (start \"\" \"%%ProgramFiles%%\\Google\\Chrome\\Application\\chrome.exe\" --app=%%URL%% & exit /b)\r\n"
           "start \"\" %%URL%%\r\n") % base_url
    url = "[InternetShortcut]\r\nURL=%s/?app=1\r\n" % base_url
    bio = io.BytesIO()
    with zipfile.ZipFile(bio, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("대흥공업 통합관리.bat", bat.encode("cp949", "replace"))
        z.writestr("대흥공업 통합관리.url", url.encode("cp949", "replace"))
        z.writestr("사용법.txt", "압축을 풀고 '대흥공업 통합관리.bat' 을 바탕화면에 두고 두 번 클릭하면 전용 창으로 열립니다.\r\n".encode("cp949"))
    return bio.getvalue()


# ──────────────────────────────────────────── 데이터 가져오기
def _norm(s):
    return re.sub(r"[\s()\[\]·./\-_]", "", str(s or "")).lower()


def _sheet_rows(ws, labels):
    """머리글 행을 찾아 {열이름: 값} 목록으로."""
    grid = [list(r) for r in ws.iter_rows(values_only=True)]
    want = {_norm(l): k for k, l in labels.items()}
    best, best_i, best_map = 0, None, {}
    for i, r in enumerate(grid[:15]):
        m = {}
        for j, v in enumerate(r):
            k = want.get(_norm(v))
            if k and k not in m:
                m[k] = j
        if len(m) > best:
            best, best_i, best_map = len(m), i, m
    if best < 2:
        return []
    out = []
    for r in grid[best_i + 1:]:
        x = {k: (r[j] if j < len(r) else None) for k, j in best_map.items()}
        if any(v not in (None, "") for v in x.values()):
            out.append(x)
    return out


def _val(v, t):
    if v is None or v == "":
        return None
    if isinstance(v, datetime.datetime):
        return v.strftime("%Y-%m-%d") if t == "date" else v.strftime("%H:%M") if t == "time" else str(v)
    if isinstance(v, datetime.time):
        return v.strftime("%H:%M")
    if t in ("num", "int"):
        try:
            f = float(str(v).replace(",", ""))
            return int(round(f)) if t == "int" else f
        except ValueError:
            return None
    if t == "yesno":
        return 1 if str(v).strip() in ("1", "Y", "y", "예", "O", "o", "True", "TRUE", "✔") else 0
    return str(v).strip()


def import_xlsm(buf, what, user, ip):
    from openpyxl import load_workbook
    backup_now("before_import")
    wb = load_workbook(io.BytesIO(buf), data_only=True, read_only=True)
    sheets = {ws.title: ws for ws in wb.worksheets}

    def find(*keys):
        for k in keys:
            for n, ws in sheets.items():
                if k in n.replace(" ", ""):
                    return ws
        return None
    plan = {"products": [("products", ("기초정보(제품)", "제품"))], "molds": [("molds", ("금형타수관리", "금형BOM")), ("parts", ("금형부품관리",))],
            "maint": [("mold_maint", ("금형정비이력", "금형부품정비이력"))], "equipment": [("equipment", ("기초정보(설비)", "설비"))],
            "production": [("production", ("생산실적통합", "생산실적")), ("defects", ("공정불량현황", "공정불량"))],
            "defects": [("defect_codes", ("기준정보(불량코드)", "불량코드"))], "workers": [("workers", ("작업자별집계", "작업자"))]}
    report = []
    for w in what:
        for tbl, keys in plan.get(w, []):
            ws = find(*[k.replace(" ", "") for k in keys])
            if not ws:
                report.append("%s: 시트를 찾지 못함 (%s)" % (meta.TABLES[tbl]["label"], ", ".join(keys)))
                continue
            spec = meta.TABLES[tbl]
            types = {c["k"]: c.get("t", "text") for c in spec["cols"]}
            rows = _sheet_rows(ws, {c["k"]: c["l"] for c in spec["cols"]})
            key = spec.get("key")
            n_new = n_upd = 0
            now = db.now()
            if tbl in ("production", "defects") and rows:
                ds = [_val(r.get("date"), "date") for r in rows if r.get("date")]
                if ds:
                    db.run('DELETE FROM "%s" WHERE date BETWEEN ? AND ?' % tbl, (min(ds), max(ds)))
            for r in rows:
                data = {k: _val(v, types[k]) for k, v in r.items()}
                if key and not data.get(key):
                    continue
                if tbl in ("production", "defects") and not data.get("date"):
                    continue
                if tbl == "mold_maint":
                    if not data.get("mold_no") or not data.get("date"):
                        continue
                    if db.one("SELECT id FROM mold_maint WHERE mold_no=? AND date=? AND COALESCE(shots,0)=?",
                              (data["mold_no"], data["date"], data.get("shots") or 0)):
                        continue
                if key and db.one('SELECT id FROM "%s" WHERE "%s"=?' % (tbl, key), (data[key],)):
                    sets = {k: v for k, v in data.items() if v is not None}
                    sets.update({"updated": now, "updated_by": user})
                    db.run('UPDATE "%s" SET %s WHERE "%s"=?' % (tbl, ",".join('"%s"=?' % k for k in sets), key), list(sets.values()) + [data[key]])
                    n_upd += 1
                else:
                    data.update({"created": now, "created_by": user, "updated": now, "updated_by": user})
                    db.run('INSERT INTO "%s" (%s) VALUES (%s)' % (tbl, ",".join('"%s"' % k for k in data), ",".join("?" * len(data))), list(data.values()))
                    n_new += 1
            report.append("%s (%s 시트): 새로 %d · 갱신 %d" % (spec["label"], ws.title, n_new, n_upd))
    db.log(user, ip, "import", "xlsm", "", "; ".join(report)[:500])
    return report


def import_partsdb(buf, user, ip):
    backup_now("before_import")
    tmp = os.path.join(db.DATA_DIR, "_import_parts.db")
    with open(tmp, "wb") as f:
        f.write(buf)
    src = sqlite3.connect(tmp)
    src.row_factory = sqlite3.Row
    try:
        tables = [r[0] for r in src.execute("SELECT name FROM sqlite_master WHERE type='table'")]
    except sqlite3.DatabaseError:
        src.close()
        os.remove(tmp)
        raise ValueError("SQLite DB 파일이 아닙니다.")
    alias = {"code": ["code", "part_code", "부품코드", "pcode"], "name": ["name", "part_name", "부품명", "pname"],
             "category": ["category", "cat", "구분", "type"], "spec": ["spec", "규격"], "material": ["material", "재질"],
             "mold_no": ["mold_no", "mold", "금형", "사용금형", "die"], "qty": ["qty", "stock", "현재고", "quantity"],
             "safety_qty": ["safety_qty", "safety", "안전재고", "min_qty"], "unit": ["unit", "단위"], "price": ["price", "단가"],
             "maker": ["maker", "vendor", "제조사", "구입처"], "location": ["location", "loc", "보관위치"], "note": ["note", "memo", "비고"]}

    def colmap(cols):
        low = {c.lower(): c for c in cols}
        return {k: next((low[a] for a in al if a in low), None) for k, al in alias.items()}
    report = []
    pt = next((t for t in tables if "part" in t.lower() and not re.search(r"tx|hist|log|inout|move", t.lower())), None)
    if not pt:
        src.close()
        os.remove(tmp)
        raise ValueError("부품 표를 찾지 못했습니다. 표: " + ", ".join(tables))
    cols = [r[1] for r in src.execute('PRAGMA table_info("%s")' % pt)]
    cm = colmap(cols)
    db.run("DELETE FROM parts_tx")
    db.run("DELETE FROM parts")
    idmap, now = {}, db.now()
    for r in src.execute('SELECT * FROM "%s"' % pt):
        r = dict(r)
        d = {k: r.get(c) for k, c in cm.items() if c}
        if not d.get("name"):
            continue
        d.update({"created": now, "created_by": user, "updated": now, "updated_by": user})
        cur = db.run('INSERT INTO parts (%s) VALUES (%s)' % (",".join('"%s"' % k for k in d), ",".join("?" * len(d))), list(d.values()))
        idmap[r.get("id") or r.get("rowid")] = cur.lastrowid
    report.append("부품 %d종 (%s 표)" % (len(idmap), pt))
    tt = next((t for t in tables if re.search(r"tx|hist|log|inout|move", t.lower())), None)
    if tt:
        tcols = [r[1].lower() for r in src.execute('PRAGMA table_info("%s")' % tt)]
        n = 0
        for r in src.execute('SELECT * FROM "%s"' % tt):
            r = {k.lower(): v for k, v in dict(r).items()}
            pid = idmap.get(r.get("part_id") or r.get("pid"))
            if not pid:
                continue
            db.run("""INSERT INTO parts_tx(part_id,date,kind,qty,after_qty,mold_no,worker,note,created,created_by)
                      VALUES(?,?,?,?,?,?,?,?,?,?)""",
                   (pid, str(r.get("date") or "")[:10], r.get("kind") or r.get("type") or "", r.get("qty"), r.get("after_qty"),
                    r.get("mold_no"), r.get("worker"), r.get("note"), now, user))
            n += 1
        report.append("입출고 이력 %d건 (%s 표)" % (n, tt))
    src.close()
    os.remove(tmp)
    db.log(user, ip, "import", "partsdb", "", "; ".join(report))
    return report
