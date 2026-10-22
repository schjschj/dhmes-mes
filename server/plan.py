# -*- coding: utf-8 -*-
"""납기 · 생산계획 — 납기현황 엑셀 → 설비별 편성 (원본 기능을 단순화해 재구성).

1) 납기 파일: 품번별 결품 수량 (+20일 / 월말 / 익월말, 누적)
2) 출하대기 재고 + 납기 파일 날짜 이후 출고분을 뺌
3) 품번 → 제품군(숫자 6자리) → 주력 설비 (금형 · 설비 표, 최근 실적에서 자동)
4) 설비마다 제품군 순서대로 근무 칸(평일 주간 / 야간 / 주말 특근)에 채움
   주간만으로 월말 결품을 못 맞추면 야간, 그래도 안 되면 케파 부족 + 주말 연장 제안
5) 원재료: 창고 재고 + 입고 예정과 필요 kg 비교
"""
import calendar
import datetime
import io
import json
import re

import db
import prodcalc
from common import get_json, set_json

DAY_MIN_DEFAULT = 560
SETUP_DEFAULT = 40
DEFAULT_SPM, DEFAULT_EFF = 300, 0.8
WD = "월화수목금토일"
WEEKEND_MIN = {"": 330, "기본": 330, "15:00": 330, "17:30": 480, "24:00": 900, "휴무": 0}
RULE_KINDS = {"night_off": "야간 안 함", "chain": "이어서 생산", "first": "늘 먼저", "stay": "옮기지 않음", "extra": "참고 노하우"}
LEARN_KINDS = {"first": "먼저 생산", "nonight": "야간 안 함"}


# ──────────────────────────────────────────── 공통
def fam_of(code):
    s = (code or "").upper().strip()
    s = re.sub(r"^H(?=\d)", "", s)
    m = re.match(r"(\d{6})", s)
    return m.group(1) if m else s.split("-")[0]


def _d(s):
    return datetime.date.fromisoformat(s)


def _month_end(d, add=0):
    y, m = d.year, d.month + add
    while m > 12:
        y, m = y + 1, m - 12
    return datetime.date(y, m, calendar.monthrange(y, m)[1])


def _holidays():
    import attendance
    return attendance.holidays_map()


def _std_minutes():
    std = prodcalc.get_std()
    base = prodcalc._m(std.get("day_start")) or 0
    sp = prodcalc._span(*std.get("work_normal"), base)
    shift = prodcalc._span(std.get("day_start"), std.get("shift_normal"), base)
    if not sp or not shift:
        return DAY_MIN_DEFAULT, 480
    brs = prodcalc._breaks(std, "normal", base)
    day = prodcalc._net((sp[0], min(sp[1], shift[1])), brs)
    night = prodcalc._net((shift[1], sp[1]), brs) if sp[1] > shift[1] else 0
    return day or DAY_MIN_DEFAULT, night


# ──────────────────────────────────────────── 납기 파일
def upload(buf, name, user):
    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(buf), data_only=True, read_only=True)
    except Exception:
        raise ValueError("엑셀 파일을 열 수 없습니다. 보안(DRM)이 걸린 파일이면 풀고 올려 주세요.")
    ws = next((w for w in wb.worksheets if "납기현황" in w.title and "제품" in w.title), None) or wb.worksheets[0]
    grid = [list(r) for r in ws.iter_rows(values_only=True)]
    hdr_i, cols = None, {}
    for i, r in enumerate(grid[:20]):
        txt = [str(x or "").replace(" ", "").replace("\n", "") for x in r]
        if any(t in ("품번", "제번", "품목코드", "PARTNO") or "품번" in t for t in txt):
            hdr_i = i
            for j, t in enumerate(txt):
                if "code" not in cols and ("품번" in t or t in ("제번", "품목코드")):
                    cols["code"] = j
                elif "name" not in cols and ("품명" in t or "품목명" in t):
                    cols["name"] = j
                elif "R" not in cols and "익월" in t:
                    cols["R"] = j
                elif "T" not in cols and ("20일" in t or "+20" in t):
                    cols["T"] = j
                elif "S" not in cols and "월말" in t:
                    cols["S"] = j
            break
    if hdr_i is None or "code" not in cols or not ({"T", "S", "R"} & set(cols)):
        raise ValueError("'품번'과 결품 열(+20일 / 월말 / 익월말)이 있는 머리글을 찾지 못했습니다. 시트: " + ws.title)

    def num(v):
        try:
            return float(str(v).replace(",", "")) if v not in (None, "") else 0.0
        except ValueError:
            return 0.0
    body = [r for r in grid[hdr_i + 1:] if r and cols["code"] < len(r) and r[cols["code"]]]
    neg = {k: any(num(r[cols[k]]) < 0 for r in body if cols.get(k) is not None and cols[k] < len(r)) for k in ("T", "S", "R") if k in cols}
    rows = {}
    for r in body:
        code = str(r[cols["code"]]).strip()
        x = rows.setdefault(code.upper(), {"code": code, "name": str(r[cols["name"]] or "") if "name" in cols and cols["name"] < len(r) else "",
                                           "T": 0, "S": 0, "R": 0})
        for k in ("T", "S", "R"):
            if k in cols and cols[k] < len(r):
                v = num(r[cols[k]])
                x[k] += (-v if v < 0 else 0) if neg.get(k) else max(0, v)
    rows = list(rows.values())
    for x in rows:   # 누적으로 맞춤 (+20일 ≤ 월말 ≤ 익월말)
        x["S"] = max(x["S"], x["T"])
        x["R"] = max(x["R"], x["S"])
    m = re.search(r"(20\d\d)[.\-_]?(\d\d)[.\-_]?(\d\d)", name or "")
    fd = "%s-%s-%s" % m.groups() if m else datetime.date.today().isoformat()
    now = db.now()
    db.run("INSERT INTO plan_uploads(file_date,file_time,name,rows,created,created_by) VALUES(?,?,?,?,?,?)",
           (fd, now[11:16], name, json.dumps(rows, ensure_ascii=False), now, user))
    db.run("DELETE FROM plan_ai")      # 새 납기 파일이면 AI 대화 지움
    return {"file_date": fd, "items": len(rows), "need": sum(1 for x in rows if x["R"] > 0)}


def _upload(date=None):
    if date:
        r = db.one("SELECT * FROM plan_uploads WHERE file_date=? ORDER BY id DESC LIMIT 1", (date,))
    else:
        r = db.one("SELECT * FROM plan_uploads ORDER BY file_date DESC, id DESC LIMIT 1")
    return r


def uploads():
    seen, out = set(), []
    for u in db.rows("SELECT file_date, file_time FROM plan_uploads ORDER BY file_date DESC, id DESC"):
        if u["file_date"] not in seen:
            seen.add(u["file_date"])
            out.append(u)
    return out


# ──────────────────────────────────────────── 금형 · 설비 표 (제품군)
def _rate(fam_codes, products, molds, day_min, night_min):
    for code in fam_codes:
        p = products.get(code.upper())
        if p and prodcalc._num(p.get("spm")):
            cav = prodcalc._num((molds.get(p.get("mold_no")) or {}).get("cavity")) or 1
            mult = prodcalc._num(p.get("valid_count")) or cav
            rpm = prodcalc._num(p["spm"]) * prodcalc._eff(p.get("efficiency")) * mult
            return rpm, False, "SPM %s × 효율 %s × %s" % (p["spm"], p.get("efficiency") or 1, mult)
    return DEFAULT_SPM * DEFAULT_EFF, True, "기본값 SPM %d × 효율 %.1f" % (DEFAULT_SPM, DEFAULT_EFF)


def _history():
    """최근 2년 실적: 제품군 → {설비: 가동일수}, 마지막 생산일."""
    since = (datetime.date.today() - datetime.timedelta(days=730)).isoformat()
    fam = {}
    for r in db.rows("SELECT code, equip_no, date FROM production WHERE date >= ?", (since,)):
        f = fam.setdefault(fam_of(r["code"]), {"eq": {}, "last": "", "dates": set()})
        f["eq"].setdefault(str(r["equip_no"]), set()).add(r["date"])
        f["dates"].add(r["date"])
        f["last"] = max(f["last"], r["date"])
    return fam


def families():
    products = {p["code"].upper(): p for p in db.rows("SELECT * FROM products")}
    molds = {m["mold_no"]: m for m in db.rows("SELECT * FROM molds")}
    stored = {f["fam"]: f for f in db.rows("SELECT * FROM plan_families")}
    hist = _history()
    day_min, night_min = _std_minutes()
    codes_by = {}
    for code, p in products.items():
        codes_by.setdefault(fam_of(p["code"]), []).append(p["code"])
    out = []
    for fam in sorted(set(codes_by) | set(stored) | set(hist)):
        codes = sorted(codes_by.get(fam, []))
        h = hist.get(fam, {"eq": {}, "last": "", "dates": set()})
        days = {int(eq): len(ds) for eq, ds in h["eq"].items() if str(eq).isdigit()}
        s = stored.get(fam)
        if s and s.get("main_eq"):
            main = s["main_eq"]
            alts = json.loads(s.get("alts") or "[]")
        else:
            ranked = sorted(days.items(), key=lambda x: -x[1])
            main = ranked[0][0] if ranked else None
            alts = [[e, n, 1] for e, n in ranked[1:]]
        known = {a[0] for a in alts} | {main}
        for e, n in days.items():
            if e not in known:
                alts.append([e, n, 0])
        rpm, dflt, src = _rate(codes, products, molds, day_min, night_min)
        locked = bool(s and s.get("locked"))
        rate_d = s["rate_d"] if locked and s.get("rate_d") else rpm * day_min
        rate_n = s["rate_n"] if locked and s.get("rate_n") else rpm * night_min
        if locked:
            rpm = rate_d / day_min if day_min else rpm
            src = "사람이 정한 값"
        out.append({"fam": fam, "codes": codes, "molds": sorted({products[c.upper()].get("mold_no") for c in codes if products[c.upper()].get("mold_no")}),
                    "main_eq": main, "alts": alts, "rpm": rpm, "dflt": dflt and not locked, "rate_d": round(rate_d), "rate_n": round(rate_n),
                    "rate_src": src, "locked": locked, "days": len(h["dates"]), "last": h["last"], "note": (s or {}).get("note") or ""})
    return out


def save_family(fam, b, user):
    def n(v):
        try:
            return float(v) if v not in (None, "") else None
        except ValueError:
            return None
    locked = 0 if b.get("auto", True) else 1
    db.run("""INSERT INTO plan_families(fam,main_eq,alts,rate_d,rate_n,locked,note,updated,updated_by) VALUES(?,?,?,?,?,?,?,?,?)
              ON CONFLICT(fam) DO UPDATE SET main_eq=excluded.main_eq, alts=excluded.alts, rate_d=excluded.rate_d,
              rate_n=excluded.rate_n, locked=excluded.locked, note=excluded.note, updated=excluded.updated, updated_by=excluded.updated_by""",
           (fam, int(b["main_eq"]) if str(b.get("main_eq") or "").isdigit() else None, json.dumps(b.get("alts") or []),
            n(b.get("rate_d")), n(b.get("rate_n")), locked, b.get("note") or "", db.now(), user))


def refresh_families():
    """최근 실적에서 새 설비 후보 개수 (표는 매번 실적으로 다시 계산됨)."""
    return {"candidates": sum(1 for f in families() for a in f["alts"] if len(a) >= 3 and not a[2] and a[1])}


# ──────────────────────────────────────────── 조정 · 입고 예정 · 노하우
def overrides():
    return db.rows("SELECT * FROM plan_overrides ORDER BY id")


def add_override(b, user, via=""):
    kind = b.get("kind")
    if kind not in ("move", "hours", "night", "first", "exclude"):
        raise ValueError("알 수 없는 조정입니다.")
    k1, k2 = str(b.get("k1") or ""), str(b.get("k2") or "")
    db.run("DELETE FROM plan_overrides WHERE kind=? AND UPPER(COALESCE(k1,''))=UPPER(?) AND COALESCE(k2,'')=?", (kind, k1, k2))
    db.run("INSERT INTO plan_overrides(kind,k1,k2,val,via,created,created_by) VALUES(?,?,?,?,?,?,?)",
           (kind, k1, k2, str(b.get("val") or ""), via, db.now(), user))


def incoming():
    rows = db.rows("SELECT * FROM plan_incoming WHERE done=0 ORDER BY COALESCE(date,'9999'), id")
    parts = {}
    for p in db.rows("SELECT mat_code, mat_desc FROM products WHERE mat_code IS NOT NULL AND mat_code<>''"):
        parts.setdefault(p["mat_code"], {"part_no": p["mat_code"], "material": p["mat_desc"] or "", "spec": ""})
    return {"rows": rows, "parts": list(parts.values())}


def save_incoming(b, user):
    try:
        kg = float(b.get("kg") or 0)
    except ValueError:
        raise ValueError("무게를 숫자로 넣으세요.")
    if not (b.get("part_no") or "").strip():
        raise ValueError("자재번호를 넣으세요.")
    if b.get("id"):
        db.run("UPDATE plan_incoming SET part_no=?, kg=?, date=?, memo=? WHERE id=?",
               (b["part_no"].strip(), kg, b.get("date") or None, b.get("memo") or "", b["id"]))
    else:
        db.run("INSERT INTO plan_incoming(part_no,kg,date,memo,created,created_by) VALUES(?,?,?,?,?,?)",
               (b["part_no"].strip(), kg, b.get("date") or None, b.get("memo") or "", db.now(), user))


def rules(active_only=False):
    sql = "SELECT * FROM plan_rules WHERE pending=0" + (" AND active=1" if active_only else "") + " ORDER BY id DESC"
    return db.rows(sql)


def add_rule(b, user, source="직접 입력", pending=0):
    kind = b.get("kind")
    if kind not in RULE_KINDS:
        raise ValueError("노하우 종류를 고르세요.")
    eq = int(b["eq"]) if str(b.get("eq") or "").isdigit() else None
    text = (b.get("text") or "").strip()
    if not text:
        text = {"night_off": "%s호기는 야간 안 함" % eq, "chain": "%s → %s 이어서 생산" % (b.get("fam"), b.get("fam2")),
                "first": "%s 제품군은 늘 먼저" % b.get("fam"), "stay": "%s 제품군은 %s호기에서 옮기지 않음" % (b.get("fam"), eq)}.get(kind, "")
    if kind == "extra" and not text:
        raise ValueError("참고 노하우는 설명을 넣어야 합니다.")
    cur = db.run("""INSERT INTO plan_rules(kind,eq,fam,fam2,text,source,evidence,active,pending,created,created_by)
                    VALUES(?,?,?,?,?,?,?,1,?,?,?)""",
                 (kind, eq, b.get("fam") or None, b.get("fam2") or None, text, source, b.get("evidence") or "", pending, db.now(), user))
    return cur.lastrowid


# ──────────────────────────────────────────── 배운 노하우 (승인된 작업지시서 vs 계획)
def _learn_cfg():
    c = get_json("plan_learn_cfg", {"days": 60, "min": 3, "ratio": 70})
    return {k: int(c.get(k, d)) for k, d in (("days", 60), ("min", 3), ("ratio", 70))}


def _wo_fams(date):
    row = db.one("SELECT data, draft FROM work_orders WHERE date=?", (date,))
    if not row:
        return None, None
    d = json.loads(row["data"])
    act = {}
    for side in ("left", "right"):
        for ln in d.get(side, []):
            c = ln["cells"]
            codes = re.findall(r"[0-9A-Z]{3,}-[0-9A-Z]+", ((c.get("code") or {}).get("t") or "") + "\n" + ((c.get("cur") or {}).get("t") or ""), re.I)
            act[str(ln.get("equip"))] = (fam_of(codes[0]) if codes else None, bool(ln.get("star")))
    return act, json.loads(row["draft"]) if row["draft"] else {}


def learn_from_wo(date):
    act, draft = _wo_fams(date)
    if act is None:
        return
    try:
        p = compute(start=date, learned=False)
    except Exception:
        return
    if p.get("none"):
        return
    obs = [o for o in get_json("plan_obs", []) if o["date"] != date]
    for e in p["eqs"]:
        eq = str(e["eq"])
        base = None
        for s in e.get("segs", []):
            if s["start"][0] <= date <= s["end"][0]:
                base = s["fam"]
                break
        base_night = any(sl == [date, "N"] for sl in e.get("night_slots", []))
        a = act.get(eq, (None, False))
        if base is None and a[0] is None:
            continue
        obs.append({"date": date, "eq": int(e["eq"]), "base_fam": base, "draft_fam": (draft or {}).get(eq),
                    "act_fam": a[0], "base_night": base_night, "act_night": a[1],
                    "ctx": {"T": {base: 0} if base else {}, "mat_short": []}})
    set_json("plan_obs", obs[-3000:])


def learned_list():
    c = _learn_cfg()
    since = (datetime.date.today() - datetime.timedelta(days=c["days"])).isoformat()
    obs = [o for o in get_json("plan_obs", []) if o["date"] >= since]
    off = set(get_json("plan_learn_off", []))
    out = {}
    by_eq = {}
    for o in obs:
        by_eq.setdefault(o["eq"], []).append(o)
    for eq, os_ in by_eq.items():
        opp = len(os_)
        diff = {}
        for o in os_:
            if o["act_fam"] and o["act_fam"] != o["base_fam"]:
                diff[o["act_fam"]] = diff.get(o["act_fam"], 0) + 1
        for fam, hits in diff.items():
            key = "first|%s|%s" % (eq, fam)
            out[key] = {"key": key, "kind": "first", "eq": eq, "fam": fam, "text": "%s호기는 %s 제품군을 먼저" % (eq, fam),
                        "hits": hits, "opp": opp, "ratio": hits * 100.0 / opp, "evidence": "최근 %d일 승인 작업지시서" % c["days"]}
        nn = sum(1 for o in os_ if o["base_night"] and not o["act_night"])
        nb = sum(1 for o in os_ if o["base_night"])
        if nn:
            key = "nonight|%s" % eq
            out[key] = {"key": key, "kind": "nonight", "eq": eq, "text": "%s호기 야간 안 함" % eq, "hits": nn, "opp": nb,
                        "ratio": nn * 100.0 / nb if nb else 0, "evidence": "계획은 야간인데 실제는 주간만"}
    for x in out.values():
        x["on_"] = x["hits"] >= c["min"] and x["ratio"] >= c["ratio"]
        x["user_off"] = x["key"] in off
    return sorted(out.values(), key=lambda x: (-x["on_"], -x["ratio"])), obs


def learned_view():
    lst, obs = learned_list()
    c = _learn_cfg()
    n = len(obs)
    def pct(k):
        xs = [o for o in obs if k(o) is not None]
        return round(sum(1 for o in xs if k(o)) * 100 / len(xs)) if xs else None
    return {"cfg": c, "learned": lst, "obs": list(reversed(obs))[:500], "rules": rules(), "kinds": LEARN_KINDS, "rule_kinds": RULE_KINDS,
            "accuracy": {"days": len({o["date"] for o in obs}), "base": pct(lambda o: o["base_fam"] == o["act_fam"] if o["act_fam"] else None),
                         "draft": pct(lambda o: (o["draft_fam"] == o["act_fam"]) if o.get("draft_fam") else None),
                         "draft_lines": sum(1 for o in obs if o.get("draft_fam")),
                         "night": pct(lambda o: o["base_night"] == o["act_night"]) if n else None}}


def learned_update(b):
    if "cfg" in b:
        c = _learn_cfg()
        for k in ("days", "min", "ratio"):
            try:
                c[k] = int(float(b["cfg"].get(k, c[k])))
            except (TypeError, ValueError):
                pass
        set_json("plan_learn_cfg", c)
    if "key" in b:
        off = set(get_json("plan_learn_off", []))
        (off.add if b.get("off") else off.discard)(b["key"])
        set_json("plan_learn_off", sorted(off))


def learned_rebuild():
    before = {x["key"]: x["on_"] for x in learned_list()[0]}
    set_json("plan_obs", [])
    since = (datetime.date.today() - datetime.timedelta(days=_learn_cfg()["days"])).isoformat()
    days = [r["date"] for r in db.rows("SELECT date FROM work_orders WHERE status='승인' AND date>=? ORDER BY date", (since,))]
    for d in days:
        learn_from_wo(d)
    after = {x["key"]: x["on_"] for x in learned_list()[0]}
    return {"days": len(days), "changed": [k for k in set(before) | set(after) if before.get(k) != after.get(k)]}


# ──────────────────────────────────────────── 계산
def _deduct(file_date):
    """출하대기 재고(불량선별 대기 보류는 빼지 않음) + 납기 파일 이후 출고 → 품번별 EA."""
    import shipping
    st = shipping.stock()
    if not st["start"]:
        return {}, False
    out = {}
    for r in st["rows"]:
        ea = max(0.0, r["avail"]) * (r["reel_qty"] or 0)
        out[r["code"].upper()] = out.get(r["code"].upper(), 0) + ea
    prods = {p["code"].upper(): p for p in db.rows("SELECT code, reel_qty FROM products")}
    for t in db.rows("SELECT code, SUM(reels) r FROM ship_tx WHERE kind='out' AND ship_date > ? GROUP BY code", (file_date,)):
        out[t["code"].upper()] = out.get(t["code"].upper(), 0) + (t["r"] or 0) * prodcalc._num((prods.get(t["code"].upper()) or {}).get("reel_qty"))
    return out, True


def _slots(start, end, eq, ov_hours, holidays, day_min, night_min, night_on):
    """[(date, 'D'|'N', 분)]"""
    out = []
    d = start
    while d <= end:
        ds = d.isoformat()
        val = ov_hours.get((str(eq), ds), ov_hours.get(("", ds)))
        off = d.weekday() >= 5 or ds in holidays
        if off:
            mins = WEEKEND_MIN.get(val or "", 0) if val is not None else 0
            if mins:
                out.append((ds, "D", mins))
        elif val != "휴무":
            out.append((ds, "D", day_min))
            if night_on and night_min:
                out.append((ds, "N", night_min))
        d += datetime.timedelta(days=1)
    return out


def _schedule(queue, slots, setup, rate_of):
    """queue: [(fam, qty)] → segs, 남은 수량, 쓴 야간 칸."""
    segs, left, nights = [], {f: q for f, q in queue}, []
    qi, prev = 0, None
    si, used = 0, 0.0
    while qi < len(queue) and si < len(slots):
        fam, _ = queue[qi]
        if left[fam] <= 0.5:
            qi += 1
            continue
        ds, kind, cap = slots[si]
        avail = cap - used
        if fam != prev:
            if kind == "N" or avail < setup:     # 교체는 주간에만
                si, used = si + 1, 0.0
                continue
            used += setup
            avail -= setup
            segs.append({"fam": fam, "qty": 0.0, "start": [ds, kind], "end": [ds, kind], "chg": [ds] if prev else None})
            prev = fam
        rate = rate_of(fam, kind)
        can = avail * rate
        take = min(can, left[fam])
        left[fam] -= take
        segs[-1]["qty"] += take
        segs[-1]["end"] = [ds, kind]
        if kind == "N" and [ds, "N"] not in nights:
            nights.append([ds, "N"])
        if take >= can - 1e-6:
            si, used = si + 1, 0.0
        else:
            used += take / rate if rate else avail
            qi += 1
    return segs, left, nights


def compute(start=None, date=None, learned=True):
    up = _upload(date)
    ai = ai_public()
    if not up:
        return {"none": True, "uploads": [], "ai": ai, "setup_default": SETUP_DEFAULT}
    today = datetime.date.today()
    if not start:
        last_ok = db.one("SELECT MAX(date) d FROM work_orders WHERE status='승인'")["d"]
        start = max(today, _d(last_ok) + datetime.timedelta(days=1)) if last_ok and last_ok >= today.isoformat() else today
        start = start.isoformat()
    s0 = _d(start)
    month_end, end = _month_end(s0), _month_end(s0, 1)
    t_date = s0 + datetime.timedelta(days=20)
    holidays = _holidays()
    day_min, night_min = _std_minutes()
    fams = {f["fam"]: f for f in families()}
    eqinfo = {str(e["no"]): {"maker": e.get("maker"), "ton": e.get("ton")} for e in db.rows("SELECT * FROM equipment")}
    ovs = overrides()
    ov = {}
    for o in ovs:
        ov.setdefault(o["kind"], {})[(o["k1"] or "").upper() if o["kind"] != "hours" else (o["k1"] or "", o["k2"])] = o["val"]
    hours = ov.get("hours", {})
    deduct, stock_used = _deduct(up["file_date"])
    products = {p["code"].upper(): p for p in db.rows("SELECT * FROM products")}
    # 규칙 · 노하우
    rule_first = {r["fam"] for r in rules(True) if r["kind"] == "first" and r["fam"]}
    rule_nonight = {r["eq"] for r in rules(True) if r["kind"] == "night_off" and r["eq"]}
    rule_stay = {(r["fam"], r["eq"]) for r in rules(True) if r["kind"] == "stay"}
    learned_on = [x for x in learned_list()[0] if x["on_"] and not x["user_off"]] if learned else []
    # 품번 → 수요
    rows, need_by_fam, unassigned = [], {}, []
    for x in json.loads(up["rows"]):
        code = x["code"]
        p = products.get(code.upper()) or {}
        ded = deduct.get(code.upper(), 0.0)
        T, S, R = (max(0.0, x[k] - ded) for k in ("T", "S", "R"))
        fam = fam_of(code)
        f = fams.get(fam) or {}
        eq = f.get("main_eq")
        moved = None
        if ov.get("move", {}).get(fam.upper()):
            eq, moved = int(ov["move"][fam.upper()]), "옮김"
        excluded = code.upper() in ov.get("exclude", {})
        r = {"code": code, "name": x.get("name") or p.get("name") or "", "fam": fam, "eq": eq, "eqs": [eq] if eq else [],
             "T": round(T), "S": round(S), "R": round(R), "ded": round(ded), "plan": 0, "plan_rl": None, "plan_box": None,
             "tier": "", "start": None, "end": None, "short": 0, "mat_short": 0, "plan_kg": None, "note": "",
             "excluded": excluded, "moved": moved}
        if not eq and R > 0:
            r["note"] = "설비 모름 (금형 · 설비 표)"
            unassigned.append(code)
        rows.append(r)
        if eq and not excluded and R > 0:
            n = need_by_fam.setdefault((eq, fam), {"T": 0, "S": 0, "R": 0, "rows": []})
            for k, v in (("T", T), ("S", S), ("R", R)):
                n[k] += v
            n["rows"].append(r)
    # 설비별 편성
    eqs = []
    learned_eq = {}
    for x in learned_on:
        learned_eq.setdefault(x["eq"], []).append(x)
    for eq in sorted({k[0] for k in need_by_fam} | {int(k) for k in eqinfo if k.isdigit()}):
        items = [(fam, n) for (e, fam), n in need_by_fam.items() if e == eq]
        e = {"eq": eq, "mode": "없음", "segs": [], "short": [], "ext": {}, "rates": {}, "changes": [], "setup": SETUP_DEFAULT,
             "need_T": 0, "need_S": 0, "need_R": 0, "cap_day": 0, "nights": 0, "learned": [], "night_slots": []}
        if not items:
            eqs.append(e)
            continue
        first = {f for f in rule_first} | {x["fam"] for x in learned_eq.get(eq, []) if x["kind"] == "first"}
        if ov.get("first"):
            first |= {k for k in ov["first"]}
        items.sort(key=lambda it: (0 if it[0].upper() in {f.upper() for f in first} else 1, -it[1]["T"], -it[1]["S"], -it[1]["R"]))
        e["learned"] = [x["text"] for x in learned_eq.get(eq, [])]

        def rate_of(fam, kind):
            f = fams.get(fam) or {}
            return ((f.get("rate_n") if kind == "N" else f.get("rate_d")) or 0) / ((night_min if kind == "N" else day_min) or 1) or DEFAULT_SPM * DEFAULT_EFF
        for fam, _ in items:
            f = fams.get(fam) or {}
            e["rates"][fam] = {"rpm": f.get("rpm"), "src": f.get("rate_src", "")}
            rd = f.get("rate_d") or DEFAULT_SPM * DEFAULT_EFF * day_min
            e["need_T"] += need_by_fam[(eq, fam)]["T"] / rd
            e["need_S"] += need_by_fam[(eq, fam)]["S"] / rd
            e["need_R"] += need_by_fam[(eq, fam)]["R"] / rd
        night_pref = ov.get("night", {}).get(str(eq).upper())
        no_night = night_pref == "off" or eq in rule_nonight or any(x["kind"] == "nonight" for x in learned_eq.get(eq, []))
        queue = [(fam, n["R"]) for fam, n in items]
        target_S = {fam: n["S"] for fam, n in items}

        def run(night_on):
            sl = _slots(s0, end, eq, hours, holidays, day_min, night_min, night_on)
            segs, left, nights = _schedule(queue, sl, SETUP_DEFAULT, rate_of)
            # 월말까지 만든 양
            made_by_me = {}
            for sg in segs:
                if sg["end"][0] <= month_end.isoformat():
                    made_by_me[sg["fam"]] = made_by_me.get(sg["fam"], 0) + sg["qty"]
                elif sg["start"][0] <= month_end.isoformat():
                    made_by_me[sg["fam"]] = made_by_me.get(sg["fam"], 0) + sg["qty"] * 0.5
            short = {f: max(0, target_S[f] - made_by_me.get(f, 0)) for f in target_S}
            return segs, left, nights, short
        e["cap_day"] = len([1 for s in _slots(s0, month_end, eq, hours, holidays, day_min, night_min, False) if s[1] == "D"])
        segs, left, nights, short = run(night_pref == "on")
        mode = "주·야" if night_pref == "on" else "주간만"
        if any(v > 0.5 for v in short.values()) and not no_night and night_pref != "on":
            segs, left, nights, short = run(True)
            mode = "주·야"
        if any(v > 0.5 for v in short.values()):
            mode = "케파 부족"
            for ds in sorted(h for h in [(s0 + datetime.timedelta(days=i)).isoformat() for i in range((month_end - s0).days + 1)]
                             if _d(h).weekday() >= 5 or h in holidays):
                if (str(eq), ds) not in hours and ("", ds) not in hours:
                    e["ext"][ds] = "24:00"
        e.update({"mode": mode, "segs": [dict(sg, qty=round(sg["qty"])) for sg in segs], "nights": len(nights), "night_slots": nights,
                  "changes": [sg["chg"][0] for sg in segs if sg.get("chg")],
                  "short": [{"fam": f, "qty": round(v)} for f, v in short.items() if v > 0.5]})
        # 품번별 결과 (제품군 생산량을 +20일 → 월말 → 익월말 순서로 품번에 나눔)
        made = {}
        for sg in segs:
            m = made.setdefault(sg["fam"], {"qty": 0, "start": sg["start"][0], "end": sg["end"][0]})
            m["qty"] += sg["qty"]
            m["end"] = sg["end"][0]
        for fam, n in items:
            m = made.get(fam, {"qty": 0, "start": None, "end": None})
            q = m["qty"]
            for tier in ("T", "S", "R"):
                for r in sorted(n["rows"], key=lambda r: -r[tier]):
                    want = r[tier] - r["plan"]
                    if want <= 0 or q <= 0:
                        continue
                    give = min(want, q)
                    r["plan"] += give
                    q -= give
            for r in n["rows"]:
                p = products.get(r["code"].upper()) or {}
                r["plan"] = round(r["plan"])
                r["start"], r["end"] = m["start"], m["end"]
                r["tier"] = "익월말" if r["plan"] >= r["R"] - 0.5 and r["R"] else "월말" if r["plan"] >= r["S"] - 0.5 and r["S"] else "+20일" if r["plan"] >= r["T"] - 0.5 and r["T"] else ""
                r["short"] = max(0, round(r["S"] - r["plan"]))
                rq = prodcalc._num(p.get("reel_qty"))
                r["plan_rl"] = round(r["plan"] / rq, 1) if rq else None
                br = prodcalc._num(p.get("box_reels"))
                r["plan_box"] = round(r["plan_rl"] / br, 1) if (r["plan_rl"] and br) else None
                w = prodcalc._num(p.get("prod_weight")) or (prodcalc._num(p.get("weight")) + prodcalc._num(p.get("scrap_weight")))
                r["plan_kg"] = round(r["plan"] * w / 1000.0, 1) if w else None
                if r["short"]:
                    r["note"] = (r["note"] + " " if r["note"] else "") + "월말 부족"
        eqs.append(e)
    # 주말 · 공휴일 날짜
    weekend_dates = []
    d = s0
    while d <= end:
        if d.weekday() >= 5 or d.isoformat() in holidays:
            weekend_dates.append(d.isoformat())
        d += datetime.timedelta(days=1)
    wk_changes = {}
    for e in eqs:
        for ds in e["changes"]:
            if ds in weekend_dates:
                wk_changes.setdefault(ds, []).append(e["eq"])
    return {"file_date": up["file_date"], "file_time": up["file_time"], "uploads": uploads(), "start": start, "end": end.isoformat(),
            "default_start": start, "eqs": eqs, "rows": rows, "overrides": ovs, "unassigned": unassigned,
            "moved": [{"fam": o["k1"], "src": (fams.get(o["k1"]) or {}).get("main_eq"), "dst": int(o["val"]), "part": False}
                      for o in ovs if o["kind"] == "move" and str(o["val"]).isdigit()],
            "eqinfo": eqinfo, "deduct_info": {"stock_used": stock_used}, "weekend_dates": weekend_dates,
            "holidays": {k: v for k, v in holidays.items() if k in weekend_dates},
            "ext": [{"eq": e["eq"], "date": ds, "mode": m} for e in eqs for ds, m in e["ext"].items()],
            "wk_changes": wk_changes, "learned": [{"text": x["text"]} for x in learned_on], "learned_skipped": [],
            "material": material(rows), "ai": ai, "setup_default": SETUP_DEFAULT,
            "t_date": t_date.isoformat(), "month_end": month_end.isoformat()}


def material(rows):
    import rawmat
    stock = rawmat.stock_by_part()
    if not stock:
        return {"ok": False}
    products = {p["code"].upper(): p for p in db.rows("SELECT * FROM products")}
    inc = db.rows("SELECT * FROM plan_incoming WHERE done=0")
    agg, noweight, nomat = {}, [], []
    for r in rows:
        if not (r["R"] or r["plan"]):
            continue
        p = products.get(r["code"].upper()) or {}
        w = prodcalc._num(p.get("prod_weight")) or (prodcalc._num(p.get("weight")) + prodcalc._num(p.get("scrap_weight")))
        if not p.get("mat_code"):
            nomat.append(r["code"])
            continue
        if not w:
            noweight.append(r["code"])
            continue
        a = agg.setdefault(p["mat_code"], {"mat": p["mat_code"], "material": p.get("mat_desc") or "", "spec": "", "codes": [],
                                           "need_T": 0, "need_S": 0, "need_R": 0, "plan": 0})
        a["codes"].append(r["code"])
        for k in ("T", "S", "R"):
            a["need_" + k] += r[k] * w / 1000.0
        a["plan"] += r["plan"] * w / 1000.0
    out = []
    for mat, a in agg.items():
        rolls = sorted(stock.get(mat, []), key=lambda x: str(x.get("isd") or ""))
        st = sum(x.get("weight") or 0 for x in rolls)
        incoming_kg = sum(i["kg"] or 0 for i in inc if i["part_no"] == mat and i["date"])
        undated = sum(i["kg"] or 0 for i in inc if i["part_no"] == mat and not i["date"])
        have = st + incoming_kg
        use, acc = [], 0
        for x in rolls:
            if acc >= a["plan"]:
                break
            use.append({"barcode": x.get("barcode"), "kg": x.get("weight") or 0, "lctn": x.get("lctn"), "isd": str(x.get("isd") or "")[:10]})
            acc += x.get("weight") or 0
        if rolls:
            a["material"] = a["material"] or rolls[0].get("material") or ""
            a["spec"] = rolls[0].get("spec") or ""
        a.update({"stock": st, "rolls": len(rolls), "incoming": incoming_kg, "undated": undated,
                  "short_S": max(0, a["need_S"] - have), "short_R": max(0, a["need_R"] - have), "use": use})
        out.append(a)
    for a in out:   # 원재료가 모자라면 품번별 부족 EA 표시
        if a["short_S"] > 0:
            for r in rows:
                if r["code"] in a["codes"]:
                    w = prodcalc._num((products.get(r["code"].upper()) or {}).get("prod_weight")) or 1
                    r["mat_short"] = round(min(r["plan"], a["short_S"] * 1000 / w / max(1, len(a["codes"]))))
    out.sort(key=lambda a: -a["short_S"])
    return {"ok": True, "rows": out, "noweight": noweight, "nomat": nomat}


# ──────────────────────────────────────────── 작업지시서 초안
def wo_draft(b, user):
    import workorders
    d1 = _d(b["date"])
    d2 = _d(b.get("date_to") or b["date"])
    if (d2 - d1).days > 6:
        d2 = d1 + datetime.timedelta(days=6)
    p = compute(start=d1.isoformat())
    if p.get("none"):
        raise ValueError("납기 파일이 없습니다.")
    dates = [(d1 + datetime.timedelta(days=i)).isoformat() for i in range((d2 - d1).days + 1)]
    exists = [ds for ds in dates if (db.one("SELECT status FROM work_orders WHERE date=?", (ds,)) or {}).get("status") == "작성"]
    if exists and not b.get("overwrite"):
        return {"exists": exists}
    made, skipped = [], []
    codes_of = {}
    for r in p["rows"]:
        if r["plan"] > 0:
            codes_of.setdefault(r["fam"], []).append(r["code"])
    for ds in dates:
        row = db.one("SELECT * FROM work_orders WHERE date=?", (ds,))
        if row and row["status"] in ("검토", "승인"):
            skipped.append({"date": ds, "why": row["status"] + "됨"})
            continue
        if not row:
            prev = db.one("SELECT date FROM work_orders WHERE date<? ORDER BY date DESC LIMIT 1", (ds,))
            workorders.create(ds, prev["date"] if prev else None, "next" if prev else "", user)
            row = db.one("SELECT * FROM work_orders WHERE date=?", (ds,))
        data = json.loads(row["data"])
        lines = {str(ln.get("equip")): ln for side in ("left", "right") for ln in data.get(side, [])}
        touched, draft = [], {}
        for e in p["eqs"]:
            ln = lines.get(str(e["eq"]))
            if not ln:
                continue
            segs = e["segs"]
            cur_codes = re.findall(r"[0-9A-Z]{3,}-[0-9A-Z]+", (ln["cells"].get("cur") or {}).get("t") or "", re.I)
            cur_fam = fam_of(cur_codes[0]) if cur_codes else None
            for i, sg in enumerate(segs):
                # 첫 제품군이 지금 걸린 제품과 다르면 계획 시작일에 교체로 봄
                first_change = i == 0 and not sg.get("chg") and sg["start"][0] == ds and cur_fam != sg["fam"]
                if (sg.get("chg") and sg["chg"][0] == ds) or first_change:
                    ln["cells"]["code"] = {"t": "\n".join(codes_of.get(sg["fam"], [sg["fam"]])[:3]), "fill": "FFFF00", "font": ""}
                    nxt = segs[i + 1] if i + 1 < len(segs) else None
                    if nxt:
                        ln["cells"]["next"] = {"t": "\n".join(codes_of.get(nxt["fam"], [nxt["fam"]])[:3]), "fill": "92D050", "font": ""}
                    touched.append(e["eq"])
                if sg["start"][0] <= ds <= sg["end"][0]:
                    draft[str(e["eq"])] = sg["fam"]
            ln["star"] = [ds, "N"] in e.get("night_slots", [])
        db.run("UPDATE work_orders SET data=?, version=version+1, draft=?, updated=?, updated_by=? WHERE id=?",
               (json.dumps(data, ensure_ascii=False), json.dumps(draft), db.now(), user, row["id"]))
        made.append({"date": ds, "touched": touched})
    return {"made": made, "skipped": skipped}


# ──────────────────────────────────────────── AI 편성 도우미 (Claude)
def _ai_cfg():
    return get_json("plan_ai_cfg", {"on": False, "key": "", "model": "claude-opus-5-5"})


def ai_public():
    c = _ai_cfg()
    return {"on": bool(c.get("on")), "key_set": bool(c.get("key")), "model": c.get("model") or "claude-opus-5-5"}


def ai_config(b):
    c = _ai_cfg()
    if b.get("clear_key"):
        c["key"] = ""
    elif b.get("key"):
        c["key"] = b["key"].strip()
    if b.get("model"):
        c["model"] = b["model"].strip()
    c["on"] = bool(b.get("on")) and bool(c.get("key"))
    set_json("plan_ai_cfg", c)


def ai_state(user_id):
    r = db.one("SELECT history FROM plan_ai WHERE user_id=?", (user_id,))
    pend = db.rows("SELECT id, kind, text, source, evidence FROM plan_rules WHERE pending=1 ORDER BY id")
    return {"ai": ai_public(), "history": json.loads(r["history"]) if r else [], "pending": pend}


AI_TOOLS = [
    {"name": "set_override", "description": "생산계획 조정을 건다. kind: move(제품군 k1 을 설비 val 로 옮김), "
     "hours(설비 k1(비우면 모든 설비) 날짜 k2 근무를 val='17:30'|'24:00'|'휴무'|'기본'), night(설비 k1 야간 val='on'|'off'), "
     "first(제품군 k1 먼저), exclude(품번 k1 계획에서 뺌).",
     "input_schema": {"type": "object", "properties": {"kind": {"type": "string", "enum": ["move", "hours", "night", "first", "exclude"]},
                                                     "k1": {"type": "string"}, "k2": {"type": "string"}, "val": {"type": "string"}},
                      "required": ["kind", "k1"]}},
    {"name": "remove_override", "description": "걸려 있는 조정을 지운다 (id 는 계획 자료의 overrides 에 있음).",
     "input_schema": {"type": "object", "properties": {"id": {"type": "integer"}}, "required": ["id"]}},
    {"name": "propose_rule", "description": "사용자가 '앞으로 …'라고 말한 편성 노하우를 제안한다 (사람이 승인해야 적용). "
     "kind: night_off(설비 eq 야간 안 함), chain(제품군 fam 다음 fam2), first(제품군 fam 늘 먼저), stay(제품군 fam 은 설비 eq 고정), extra(설명만).",
     "input_schema": {"type": "object", "properties": {"kind": {"type": "string", "enum": list(RULE_KINDS)}, "eq": {"type": "integer"},
                                                     "fam": {"type": "string"}, "fam2": {"type": "string"}, "text": {"type": "string"}},
                      "required": ["kind", "text"]}},
    {"name": "fill_workorder", "description": "계획대로 작업지시서(교체 제번·차기생산품·야간 ★)를 채운다. 승인·검토된 날은 건너뜀.",
     "input_schema": {"type": "object", "properties": {"date": {"type": "string"}, "date_to": {"type": "string"}}, "required": ["date"]}},
]


def _plan_brief(p):
    if p.get("none"):
        return "납기 파일 없음."
    eqs = [{"eq": e["eq"], "mode": e["mode"], "need_days": [round(e["need_T"], 1), round(e["need_S"], 1), round(e["need_R"], 1)],
            "segs": [[s["fam"], s["qty"], s["start"][0], s["end"][0]] for s in e["segs"]], "short": e["short"]}
           for e in p["eqs"] if e["mode"] != "없음"]
    rows = [[r["code"], r["fam"], r["eq"], r["T"], r["S"], r["R"], r["plan"], r["note"]] for r in p["rows"] if r["R"] or r["plan"]]
    return json.dumps({"file_date": p["file_date"], "start": p["start"], "end": p["end"], "eqs": eqs,
                       "rows[code,fam,eq,T,S,R,plan,note]": rows[:400],
                       "overrides": [{k: o[k] for k in ("id", "kind", "k1", "k2", "val")} for o in p["overrides"]]}, ensure_ascii=False)


def ai_ask(me, b):
    """Claude 에게 묻고, 도구(조정·노하우 제안·작업지시서 채우기)를 쓰게 한다."""
    if b.get("reset"):
        db.run("DELETE FROM plan_ai WHERE user_id=?", (me["id"],))
        db.run("DELETE FROM plan_rules WHERE pending=1")
        return {"ok": True}
    c = _ai_cfg()
    if not c.get("on") or not c.get("key"):
        raise ValueError("AI 설정에서 API 키를 넣고 켜 주세요.")
    try:
        import anthropic
    except ImportError:
        raise ValueError("서버 PC 에 anthropic 패키지가 없습니다. 서버 PC 에서 'pip install anthropic' 을 실행하세요.")
    st = ai_state(me["id"])
    hist = st["history"]
    p = compute()
    wo_txt = ""
    if b.get("page") == "workorder" and b.get("date"):
        import workorders
        wo_txt = "\n작업지시서 %s 요약: %s" % (b["date"], json.dumps(workorders.summary(b["date"]), ensure_ascii=False))
    know = "\n".join("- " + r["text"] for r in rules(True))
    system = ("너는 프레스 공장(대흥공업) 생산 편성 도우미다. 생산계획 · 작업지시서 자료만 보고, 한국어로 짧고 정확하게 답한다. "
              "계획을 바꿔 달라고 하면 set_override / remove_override 도구를 쓰고, '앞으로 …'처럼 계속 지킬 규칙은 propose_rule 로 제안만 한다. "
              "작업지시서를 채워 달라면 fill_workorder 를 쓴다. 숫자는 자료에 있는 것만 말하고 모르면 모른다고 한다.\n"
              "T=+20일 결품, S=월말, R=익월말 (누적 EA). 사용자: %s\n지켜야 할 노하우:\n%s" % (me["name"], know or "(없음)"))
    messages = []
    for m in hist[-20:]:
        messages.append({"role": "user" if m["role"] == "user" else "assistant", "content": m["text"]})
    messages.append({"role": "user", "content": "[현재 생산계획 자료]\n%s%s\n\n[질문]\n%s" % (_plan_brief(p), wo_txt, b.get("text", ""))})
    client = anthropic.Anthropic(api_key=c["key"])
    model = c.get("model") or "claude-opus-5-5"
    kw = {}
    if model in ("claude-opus-5-5", "claude-opus-5", "claude-sonnet-5-5", "claude-fable-5-1"):
        kw = {"betas": ["server-side-fallback-2026-07-01"], "fallbacks": "default"}   # 안전 분류기가 거절하면 서버가 다른 모델로 이어 감
    applied, wo_changed, text = [], False, ""
    for _ in range(8):
        try:
            if kw:
                resp = client.beta.messages.create(model=model, max_tokens=16000, system=system, tools=AI_TOOLS, messages=messages, **kw)
            else:
                resp = client.messages.create(model=model, max_tokens=16000, system=system, tools=AI_TOOLS, messages=messages)
        except anthropic.AuthenticationError:
            raise ValueError("API 키가 올바르지 않습니다.")
        except anthropic.RateLimitError:
            raise ValueError("AI 요청이 너무 많습니다. 잠시 뒤 다시 해 주세요.")
        except anthropic.APIConnectionError:
            raise ValueError("AI 서버에 연결하지 못했습니다. 서버 PC 인터넷을 확인하세요.")
        except anthropic.APIStatusError as e:
            raise ValueError("AI 오류 (%s): %s" % (e.status_code, e.message))
        if resp.stop_reason == "refusal":
            text = "이 요청에는 답할 수 없습니다."
            break
        text = "\n".join(bl.text for bl in resp.content if getattr(bl, "type", "") == "text").strip() or text
        if resp.stop_reason != "tool_use":
            break
        messages.append({"role": "assistant", "content": resp.content})
        results = []
        for bl in resp.content:
            if getattr(bl, "type", "") != "tool_use":
                continue
            try:
                out = _ai_tool(bl.name, dict(bl.input), me)
                if bl.name == "fill_workorder":
                    wo_changed = True
                applied.append(out)
                results.append({"type": "tool_result", "tool_use_id": bl.id, "content": out})
            except Exception as e:
                results.append({"type": "tool_result", "tool_use_id": bl.id, "content": "실패: %s" % e, "is_error": True})
        messages.append({"role": "user", "content": results})
    hist.append({"role": "user", "text": b.get("text", ""), "page": b.get("page")})
    hist.append({"role": "assistant", "text": text or "(답 없음)", "applied": applied})
    db.run("""INSERT INTO plan_ai(user_id,history,updated) VALUES(?,?,?)
              ON CONFLICT(user_id) DO UPDATE SET history=excluded.history, updated=excluded.updated""",
           (me["id"], json.dumps(hist[-60:], ensure_ascii=False), db.now()))
    return {"applied": applied, "wo_changed": wo_changed}


def _ai_tool(name, a, me):
    if name == "set_override":
        add_override(a, me["name"], via="AI")
        return "조정: %s %s %s %s" % (a.get("kind"), a.get("k1"), a.get("k2") or "", a.get("val") or "")
    if name == "remove_override":
        db.run("DELETE FROM plan_overrides WHERE id=?", (a["id"],))
        return "조정 %s 지움" % a["id"]
    if name == "propose_rule":
        add_rule(a, me["name"], source="AI 대화", pending=1)
        return "노하우 제안: %s" % a.get("text")
    if name == "fill_workorder":
        r = wo_draft({"date": a["date"], "date_to": a.get("date_to") or a["date"], "overwrite": True}, me["name"])
        return "작업지시서 채움: %s" % ", ".join(m["date"] for m in r.get("made", [])) or "없음"
    raise ValueError("알 수 없는 도구")
