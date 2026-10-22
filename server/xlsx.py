# -*- coding: utf-8 -*-
"""엑셀 내보내기.
POST /api/xlsx        {filename, sheets:[{name, head[], rows[[]], fmts?[]}]}          — 일반 목록
POST /api/xlsx_styled {filename, sheets:[{name, cells[[r,c,v,style]], merges[], widths{}, heights{},
                       freeze[rows,cols], landscape, fit, scale, col_breaks[], notes[[r,c,text,side]]}]}  — 근태 양식
"""
import io

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.pagebreak import Break

_HEAD_FILL = PatternFill("solid", fgColor="1D2733")
_HEAD_FONT = Font(color="FFFFFF", bold=True)
_THIN = Side(style="thin", color="D5DCE4")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_BLACK = Side(style="thin", color="808080")
_BORDER2 = Border(left=_BLACK, right=_BLACK, top=_BLACK, bottom=_BLACK)


def _save(wb):
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def build(payload):
    wb = Workbook()
    wb.remove(wb.active)
    for sheet in payload.get("sheets", []):
        ws = wb.create_sheet((sheet.get("name") or "Sheet")[:31])
        head = sheet.get("head", [])
        fmts = sheet.get("fmts") or []
        if head:
            ws.append(head)
            for ci, _ in enumerate(head, 1):
                cell = ws.cell(row=1, column=ci)
                cell.fill = _HEAD_FILL
                cell.font = _HEAD_FONT
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.border = _BORDER
        for r in sheet.get("rows", []):
            ws.append(["" if v is None else v for v in r])
        for ci, f in enumerate(fmts, 1):
            if f in ("pct", "pct2"):
                for row in ws.iter_rows(min_row=2, min_col=ci, max_col=ci):
                    for c in row:
                        c.number_format = "0.0%" if f == "pct" else "0.00%"
        for ci, _ in enumerate(head, 1):
            width = len(str(head[ci - 1])) + 2
            for r in sheet.get("rows", []):
                if ci - 1 < len(r) and r[ci - 1] is not None:
                    width = max(width, min(40, len(str(r[ci - 1])) + 2))
            ws.column_dimensions[get_column_letter(ci)].width = width
        ws.freeze_panes = "A2"
    if not wb.worksheets:
        wb.create_sheet("Sheet")
    return _save(wb)


def build_styled(payload):
    wb = Workbook()
    wb.remove(wb.active)
    for sh in payload.get("sheets", []):
        ws = wb.create_sheet((sh.get("name") or "Sheet")[:31])
        for r, c, v, st in sh.get("cells", []):
            st = st or {}
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = Font(bold=bool(st.get("bold")), size=st.get("size") or 10, color=st.get("font") or "000000")
            if st.get("fill"):
                cell.fill = PatternFill("solid", fgColor=st["fill"])
            cell.alignment = Alignment(horizontal=st.get("h") or "center", vertical="center", wrap_text=st.get("wrap", False))
            if st.get("border", True):
                cell.border = _BORDER2
            if st.get("num"):
                cell.number_format = "0.0"
        for m in sh.get("merges", []):
            ws.merge_cells(m)
        for col, w in (sh.get("widths") or {}).items():
            ws.column_dimensions[get_column_letter(int(col))].width = w
        for row, h in (sh.get("heights") or {}).items():
            ws.row_dimensions[int(row)].height = h
        fr = sh.get("freeze") or [0, 0]
        if fr[0] or fr[1]:
            ws.freeze_panes = ws.cell(row=fr[0] + 1, column=fr[1] + 1)
        for r, c, text, _side in sh.get("notes", []):
            cm = Comment(text, sh.get("note_author") or "")
            cm.width, cm.height = 200, 80
            ws.cell(row=r, column=c).comment = cm
        ps = ws.page_setup
        ps.orientation = "landscape" if sh.get("landscape") else "portrait"
        ps.paperSize = ws.PAPERSIZE_A4
        if sh.get("scale"):
            ps.scale = int(sh["scale"])
        elif sh.get("fit"):
            ws.sheet_properties.pageSetUpPr.fitToPage = True
            ps.fitToWidth, ps.fitToHeight = 1, 0
        for cb in sh.get("col_breaks") or []:
            if cb:
                ws.col_breaks.append(Break(id=int(cb)))
        ws.page_margins.left = ws.page_margins.right = 0.25
        ws.page_margins.top = ws.page_margins.bottom = 0.4
    if not wb.worksheets:
        wb.create_sheet("Sheet")
    return _save(wb)
