# -*- coding: utf-8 -*-
"""금형수리이력카드 가져오기 (.xlsx / .xls / .zip).
카드 양식: 시트마다 금형 하나. '금형번호' 칸 오른쪽에 번호, 아래 표 머리글(일자 · 문제점 · 원인 · 수리내용 · 수리유형 · 릴 · 수량 · 수리자)
.xls(옛 엑셀)는 xlrd 패키지가 있어야 읽음 — 없으면 '못 읽은 파일'로 알려 줌.
"""
import datetime
import hashlib
import io
import re
import zipfile

import db

HEAD = {"date": ["일자", "날짜", "수리일"], "code": ["제번", "품번"], "problem": ["문제점", "의뢰"], "cause": ["원인"],
        "fix": ["수리내용", "조치"], "types": ["수리유형", "유형"], "reels": ["릴", "r/l"], "qty": ["수량", "ea"],
        "worker": ["수리자", "작업자", "담당"], "note": ["비고"]}


def _grid_xlsx(buf):
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(buf), data_only=True, read_only=True)
    return [(ws.title, [list(r) for r in ws.iter_rows(values_only=True)]) for ws in wb.worksheets]


def _grid_xls(buf):
    import xlrd
    wb = xlrd.open_workbook(file_contents=buf)
    out = []
    for sh in wb.sheets():
        rows = []
        for r in range(sh.nrows):
            row = []
            for c in range(sh.ncols):
                cell = sh.cell(r, c)
                v = cell.value
                if cell.ctype == xlrd.XL_CELL_DATE:
                    v = datetime.datetime(*xlrd.xldate_as_tuple(v, wb.datemode))
                row.append(v)
            rows.append(row)
        out.append((sh.name, rows))
    return out


def _date(v):
    if isinstance(v, datetime.datetime):
        return v.date().isoformat()
    if isinstance(v, datetime.date):
        return v.isoformat()
    s = str(v or "").strip()
    m = re.match(r"(\d{2,4})[.\-/년 ]+(\d{1,2})[.\-/월 ]+(\d{1,2})", s)
    if not m:
        return None
    y = int(m.group(1))
    y = y + 2000 if y < 100 else y
    try:
        return datetime.date(y, int(m.group(2)), int(m.group(3))).isoformat()
    except ValueError:
        return None


def _sheet(name, title, grid, types, molds, res, user):
    mold = None
    for r in grid[:12]:
        for j, v in enumerate(r):
            if v and "금형" in str(v) and ("번호" in str(v) or "no" in str(v).lower()):
                for x in r[j + 1:]:
                    if x not in (None, ""):
                        mold = str(x).strip().upper()
                        break
            if mold:
                break
        if mold:
            break
    if not mold:
        res["skipped_sheets"].append("%s · %s (금형번호 없음)" % (name, title))
        return
    if mold not in molds:
        res["no_mold"].append(mold)
    hi, cols = None, {}
    for i, r in enumerate(grid):
        txt = [str(v or "").replace(" ", "").lower() for v in r]
        m = {}
        for k, cands in HEAD.items():
            for j, t in enumerate(txt):
                if t and any(c.lower() in t for c in cands) and k not in m and j not in m.values():
                    m[k] = j
                    break
        if "date" in m and ("fix" in m or "problem" in m):
            hi, cols = i, m
            break
    if hi is None:
        res["skipped_sheets"].append("%s · %s (표 머리글 없음)" % (name, title))
        return
    for i, r in enumerate(grid[hi + 1:], hi + 2):
        g = lambda k: r[cols[k]] if k in cols and cols[k] < len(r) else None
        if not any(g(k) not in (None, "") for k in ("problem", "fix", "cause")):
            continue
        d = _date(g("date"))
        if not d:
            res["bad_dates"].append("%s · %s %d행 '%s'" % (name, title, i, g("date")))
            res["bad"] += 1
            continue
        tv = str(g("types") or "").strip()
        tcodes = []
        for t in re.split(r"[,/ ·]+", tv):
            if not t:
                continue
            code = t.split(".")[0]
            if code in types:
                tcodes.append(code)
            else:
                hit = next((c for c, n in types.items() if n and n in t), None)
                if hit:
                    tcodes.append(hit)
                else:
                    res["odd_types"].append(t)
        key = hashlib.md5(("%s|%s|%s|%s" % (mold, d, g("problem"), g("fix"))).encode("utf-8")).hexdigest()
        if db.one("SELECT id FROM mold_repair WHERE src_key=?", (key,)):
            res["dup"] += 1
            continue

        def num(v, as_int=False):
            try:
                f = float(str(v).replace(",", ""))
                return int(round(f)) if as_int else f
            except (TypeError, ValueError):
                return None
        now = db.now()
        db.run("""INSERT INTO mold_repair(mold_no,code,date,problem,cause,fix,types,reels,qty,worker,note,src,src_key,
                  created,created_by,updated,updated_by) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
               (mold, str(g("code") or "").strip() or None, d, g("problem"), g("cause"), g("fix"), ",".join(dict.fromkeys(tcodes)),
                num(g("reels")), num(g("qty"), True), g("worker"), g("note"), "%s · %s · %d행" % (name, title, i), key, now, user, now, user))
        res["added"] += 1


def import_file(name, buf, user):
    res = {"files": 0, "added": 0, "dup": 0, "bad": 0, "fixed": [], "unread": [], "skipped_sheets": [], "no_mold": [],
           "bad_dates": [], "odd_types": []}
    files = []
    if name.lower().endswith(".zip"):
        with zipfile.ZipFile(io.BytesIO(buf)) as z:
            for n in z.namelist():
                if n.lower().endswith((".xls", ".xlsx")):
                    files.append((n.split("/")[-1], z.read(n)))
    else:
        files = [(name, buf)]
    types = {t["code"]: t["name"] for t in db.rows("SELECT code, name FROM repair_types")}
    molds = {m["mold_no"].upper() for m in db.rows("SELECT mold_no FROM molds")}
    for n, b in files:
        try:
            sheets = _grid_xlsx(b) if n.lower().endswith(".xlsx") else _grid_xls(b)
        except ImportError:
            res["unread"].append("%s (.xls 는 서버에 xlrd 설치 필요: pip install xlrd)" % n)
            continue
        except Exception as e:
            res["unread"].append("%s (%s)" % (n, e))
            continue
        res["files"] += 1
        for title, grid in sheets:
            _sheet(n, title, grid, types, molds, res, user)
    res["no_mold"] = sorted(set(res["no_mold"]))
    res["odd_types"] = sorted(set(res["odd_types"]))[:50]
    return res
